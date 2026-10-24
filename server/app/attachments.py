"""Files users attach to messages: classify, extract document text, normalize images for vision models."""

from __future__ import annotations

import io
import mimetypes
from dataclasses import dataclass

MAX_SIZE = 20 * 1024 * 1024
MAX_TEXT = 200_000
IMAGE_MAX_SIDE = 1568  # what vision models actually use; larger images only cost tokens
IMAGE_MAX_BYTES = 3_500_000
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".xml", ".html", ".htm", ".css", ".js", ".ts", ".tsx", ".jsx", ".py", ".rb", ".go", ".rs", ".java", ".kt", ".swift",
    ".c", ".h", ".cpp", ".hpp", ".cs", ".php", ".sh", ".ps1", ".bat", ".sql", ".log", ".env", ".dockerfile", ".tf",
}


class AttachmentError(Exception):
    pass


@dataclass
class Processed:
    filename: str
    mime: str
    kind: str  # image | document | binary
    data: bytes
    text: str | None


def _image(filename: str, data: bytes) -> Processed:
    from PIL import Image, UnidentifiedImageError

    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except (UnidentifiedImageError, OSError) as exc:
        raise AttachmentError(f"{filename} is not a readable image") from exc
    fmt = (img.format or "").upper()
    if fmt in ("PNG", "JPEG", "GIF", "WEBP") and max(img.size) <= IMAGE_MAX_SIDE and len(data) <= IMAGE_MAX_BYTES:
        return Processed(filename, Image.MIME[fmt], "image", data, None)
    if max(img.size) > IMAGE_MAX_SIDE:
        img.thumbnail((IMAGE_MAX_SIDE, IMAGE_MAX_SIDE))
    out = io.BytesIO()
    if img.mode in ("RGBA", "LA", "P"):
        img.save(out, "PNG", optimize=True)
        mime = "image/png"
    else:
        img.convert("RGB").save(out, "JPEG", quality=88)
        mime = "image/jpeg"
    return Processed(filename, mime, "image", out.getvalue(), None)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"--- page {i} ---\n{page.extract_text() or ''}")
        if sum(len(p) for p in pages) > MAX_TEXT:
            break
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def process(filename: str, mime: str | None, data: bytes) -> Processed:
    if len(data) > MAX_SIZE:
        raise AttachmentError(f"{filename} is larger than {MAX_SIZE // (1024 * 1024)} MB")
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1][:300] or "attachment"
    mime = (mime or mimetypes.guess_type(filename)[0] or "application/octet-stream").lower()
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if mime.startswith("image/"):
        return _image(filename, data)
    text: str | None = None
    try:
        if mime == "application/pdf" or ext == ".pdf":
            text, mime = _pdf_text(data), "application/pdf"
        elif ext == ".docx" or "wordprocessingml" in mime:
            text = _docx_text(data)
        elif mime.startswith("text/") or ext in TEXT_EXTENSIONS or "json" in mime or "xml" in mime:
            text = data.decode("utf-8", errors="replace")
        elif b"\0" not in data[:8192]:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = None
    except Exception as exc:  # noqa: BLE001 - a broken document should not break the chat
        raise AttachmentError(f"cannot read {filename}: {exc}") from exc
    if text is None:
        return Processed(filename, mime, "binary", data, None)
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT] + f"\n… (truncated, {len(text)} characters in total)"
    return Processed(filename, mime, "document", data, text)
