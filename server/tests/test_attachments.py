"""Attachments: upload, text extraction, images to vision models, history; and reusing an empty session."""

import base64
import io

import docx
from PIL import Image

from app.providers.base import ChatMessage, ImagePart
from app.providers.openai_compat import to_openai_messages
from tests.conftest import chat, finished, models_by_name

def make_pdf(text: str) -> bytes:
    """A minimal but well-formed one-page PDF (with xref table) containing `text`."""
    stream = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj" % i + body + b"endobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return bytes(out)


PDF = make_pdf("Invoice total: 42 EUR")


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


async def upload(client, sid, filename, data, mime=None):
    r = await client.post(f"/api/sessions/{sid}/attachments", json={"filename": filename, "mime": mime, "data_base64": b64(data)})
    assert r.status_code == 201, r.text
    return r.json()


async def test_documents_and_images_reach_the_models(client):
    models = await models_by_name(client)
    ids = [models["mock-alpha"]["id"], models["mock-beta"]["id"]]
    s = (await client.post("/api/sessions", json={"title": "files", "model_ids": ids})).json()

    note = await upload(client, s["id"], "notes.md", b"# Notes\nThe deadline is Friday.")
    pdf = await upload(client, s["id"], "invoice.pdf", PDF)
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("Spec: the button must be blue.")
    d.save(buf)
    word = await upload(client, s["id"], "spec.docx", buf.getvalue())
    big = io.BytesIO()
    Image.new("RGB", (3000, 2000), "orange").save(big, "PNG")
    img = await upload(client, s["id"], "screenshot.png", big.getvalue(), "image/png")
    blob = await upload(client, s["id"], "data.bin", b"\x00\x01\x02" * 100)

    assert note["kind"] == pdf["kind"] == word["kind"] == "document"
    assert pdf["text_chars"] > 0 and blob["kind"] == "binary"
    assert img["kind"] == "image" and img["mime"] in ("image/jpeg", "image/png")
    stored = Image.open(io.BytesIO((await client.get(f"/api/attachments/{img['id']}")).content))
    assert max(stored.size) == 1568  # downscaled for vision models

    events = await chat(client, s["id"], "What do these say?", attachment_ids=[note["id"], pdf["id"], word["id"], img["id"], blob["id"]])
    replies = finished(events)
    assert len(replies) == 2
    for r in replies:
        assert "(I see 1 image.)" in r["content"]
        assert "notes.md" in r["content"] and "invoice.pdf" in r["content"] and "spec.docx" in r["content"]
    user = next(e["message"] for e in events if e["type"] == "message" and e["message"]["role"] == "user")
    assert [a["filename"] for a in user["attachments"]] == ["notes.md", "invoice.pdf", "spec.docx", "screenshot.png", "data.bin"]
    history = (await client.get(f"/api/sessions/{s['id']}/messages")).json()
    assert history[0]["attachments"][3]["kind"] == "image"

    # extracted text is what the model receives
    from app import db as dbmod
    from app.agent.loop import load_attachments
    from app.tables import Message

    async with dbmod.sessionmaker()() as db:
        msg = await db.get(Message, user["id"])
        files = await load_attachments(db, [msg])
    assert "Invoice total: 42 EUR" in files[pdf["id"]].text
    assert "button must be blue" in files[word["id"]].text


async def test_attachment_rules(client):
    models = await models_by_name(client)
    a = (await client.post("/api/sessions", json={"title": "a", "model_ids": [models["mock-alpha"]["id"]]})).json()
    b = (await client.post("/api/sessions", json={"title": "b", "model_ids": [models["mock-alpha"]["id"]]})).json()
    f = await upload(client, a["id"], "x.txt", b"hello")
    r = await client.post(f"/api/sessions/{b['id']}/chat", json={"content": "hi", "attachment_ids": [f["id"]]})
    assert r.status_code == 409 and "does not belong" in r.json()["detail"]
    bad = await client.post(f"/api/sessions/{a['id']}/attachments", json={"filename": "x.png", "mime": "image/png", "data_base64": b64(b"nope")})
    assert bad.status_code == 422


def test_openai_image_parts():
    msgs = to_openai_messages("", [ChatMessage("user", "look", images=[ImagePart("image/png", "AAAA")])])
    assert msgs[0]["content"][1] == {"type": "image_url", "image_url": {"url": "data:image/png;base64,AAAA"}}


async def test_new_session_reuses_an_empty_one(client):
    models = await models_by_name(client)
    alpha, beta = models["mock-alpha"]["id"], models["mock-beta"]["id"]
    first = (await client.post("/api/sessions", json={"title": "empty", "model_ids": [alpha]})).json()
    again = (await client.post("/api/sessions", json={"title": "second try", "model_ids": [beta], "reuse_empty": True})).json()
    assert again["id"] == first["id"] and again["title"] == "second try"
    assert [p["model_id"] for p in again["participants"]] == [beta]
    await chat(client, first["id"], "now it has messages")
    fresh = (await client.post("/api/sessions", json={"title": "third", "model_ids": [beta], "reuse_empty": True})).json()
    assert fresh["id"] != first["id"]
